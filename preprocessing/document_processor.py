from pathlib import Path

import pymupdf
import pytesseract
from PIL import Image
from docx import Document
from openpyxl import load_workbook


# ---------------------------------------------------------
# TXT
# ---------------------------------------------------------

def extract_from_txt(file_path: str) -> str:
    path = Path(file_path)
    return path.read_text(encoding="utf-8")


# ---------------------------------------------------------
# PDF OCR helper
# ---------------------------------------------------------

def ocr_pdf_page(page) -> str:
    """
    Convert a PDF page into an image and run OCR.
    """

    # Render PDF page at higher resolution for better OCR
    matrix = pymupdf.Matrix(2, 2)
    pixmap = page.get_pixmap(matrix=matrix, alpha=False)

    image = Image.frombytes(
        "RGB",
        [pixmap.width, pixmap.height],
        pixmap.samples,
    )

    text = pytesseract.image_to_string(image)

    return text.strip()


# ---------------------------------------------------------
# PDF
# ---------------------------------------------------------

def extract_from_pdf(file_path: str) -> tuple[str, bool, list[int]]:
    """
    Extract text from a PDF.

    If a page contains little/no selectable text,
    OCR is automatically used for that page.

    Returns:
        text
        ocr_used
        ocr_pages
    """

    document = pymupdf.open(file_path)

    pages = []
    ocr_pages = []

    for page_number, page in enumerate(document, start=1):

        # First try normal PDF text extraction
        text = page.get_text("text").strip()

        # If the page has little/no text, use OCR
        if len(text) < 30:
            print(f"  OCR processing page {page_number}...")
            text = ocr_pdf_page(page)

            if text:
                ocr_pages.append(page_number)

        if text:
            pages.append(text)

    document.close()

    return (
        "\n\n".join(pages).strip(),
        len(ocr_pages) > 0,
        ocr_pages,
    )


# ---------------------------------------------------------
# DOCX
# ---------------------------------------------------------

def extract_from_docx(file_path: str) -> str:
    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    return "\n".join(paragraphs).strip()


# ---------------------------------------------------------
# XLSX
# ---------------------------------------------------------

def extract_from_xlsx(file_path: str) -> str:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    sheets = []

    for worksheet in workbook.worksheets:

        sheets.append(f"[Worksheet: {worksheet.title}]")

        for row in worksheet.iter_rows(values_only=True):

            values = []

            for cell in row:
                if cell is not None:
                    values.append(str(cell))

            if values:
                sheets.append("\t".join(values))

    workbook.close()

    return "\n".join(sheets).strip()


# ---------------------------------------------------------
# Unified document extraction
# ---------------------------------------------------------

def extract_text(file_path: str) -> dict:

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = path.suffix.lower()

    ocr_used = False
    ocr_pages = []

    # ---------------------------------------------
    # TXT
    # ---------------------------------------------

    if extension == ".txt":

        text = extract_from_txt(file_path)

    # ---------------------------------------------
    # PDF
    # ---------------------------------------------

    elif extension == ".pdf":

        text, ocr_used, ocr_pages = extract_from_pdf(
            file_path
        )

    # ---------------------------------------------
    # DOCX
    # ---------------------------------------------

    elif extension == ".docx":

        text = extract_from_docx(file_path)

    # ---------------------------------------------
    # XLSX
    # ---------------------------------------------

    elif extension == ".xlsx":

        text = extract_from_xlsx(file_path)

    # ---------------------------------------------
    # Unsupported
    # ---------------------------------------------

    else:

        raise ValueError(
            f"Unsupported file type: {extension}. "
            "Supported types: .txt, .pdf, .docx, .xlsx"
        )

    return {
        "file_name": path.name,
        "file_type": extension,
        "text": text,
        "character_count": len(text),
        "ocr_used": ocr_used,
        "ocr_pages": ocr_pages,
    }


# ---------------------------------------------------------
# Test
# ---------------------------------------------------------

if __name__ == "__main__":

    import sys

    if len(sys.argv) < 2:
        print("Usage:")
        print(
            "python -m preprocessing.document_processor <file_path>"
        )
        sys.exit(1)

    file_path = sys.argv[1]

    result = extract_text(file_path)

    print("\nDocument Processing Result")
    print("=" * 60)

    print(f"File       : {result['file_name']}")
    print(f"Type       : {result['file_type']}")
    print(f"Characters : {result['character_count']}")
    print(f"OCR Used   : {result['ocr_used']}")
    print(f"OCR Pages  : {result['ocr_pages']}")

    print("\nExtracted Text")
    print("-" * 60)
    print(result["text"])
    print("-" * 60)